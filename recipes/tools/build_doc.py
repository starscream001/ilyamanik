import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
DATA=json.loads((ROOT/'data/recipes.json').read_text())
R=DATA['recipes']; N=DATA['notes']
OUT=ROOT/'downloads';OUT.mkdir(exist_ok=True)
doc=Document(); sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.6);sec.bottom_margin=Cm(1.6);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
for sty in ['Normal','Body Text','List Number','List Bullet']:
 s=doc.styles[sty];s.font.name='Arial';s.font.size=Pt(10);s.paragraph_format.space_after=Pt(5);s.paragraph_format.line_spacing=1.12
for sty,size in [('Title',30),('Heading 1',21),('Heading 2',12)]:
 s=doc.styles[sty];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0);s.paragraph_format.space_after=Pt(8)
footer=sec.footer.paragraphs[0];footer.alignment=2
run=footer.add_run('Книга рецептов  ·  ');run.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def p(text,style=None):return doc.add_paragraph(text,style)
def head(text):doc.add_heading(text,2)
def table(rows,headers):
 t=doc.add_table(rows=1, cols=len(headers)); t.style='Normal Table'
 for cell,text in zip(t.rows[0].cells,headers):cell.text=text
 trPr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for cell,text in zip(cells,row):cell.text=str(text)
 for row in t.rows:
  pr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');pr.append(ns)
  for cell in row.cells:
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1
    for run in para.runs:run.font.size=Pt(9)
 # neutral shading
 for cell in t.rows[0].cells:
  tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'EBE8E1');tcPr.append(shd)
 return t
for style in doc.styles:
 if style.type==1:
  for border in list(style.element.xpath('./w:pPr/w:pBdr')):border.getparent().remove(border)
p('Книга рецептов','Title')
p('Рецепты и кухонные записи из четырёх сканированных тетрадей')
p('В книге собраны составы блюд, способы приготовления, варианты и рукописные дополнения. Рецепты расположены по разделам: от завтраков и супов до выпечки и напитков. Технологические карты вынесены в отдельный раздел и упорядочены по исходным номерам.')
p('Всего 73 рецептурные записи и 15 групп заметок. Повторные сканы объединены, продолжения присоединены к соответствующим рецептам. Неполные записи обозначены как фрагменты. Противоречащие друг другу цифры и неуверенно прочитанные слова приведены в примечаниях к конкретному рецепту.')
head('Как читать количества')
p('Пары значений вида 800/670 г означают брутто/нетто: массу до и после обработки. Если в источнике единица не записана или исправление неразборчиво, это указано явно. Объём стакана и масса пачки не подменены предполагаемыми значениями. Выходы сохранены по оригиналу и могут не совпадать с суммой закладки.')
p('Ссылки вида «PDF 3, с. 12» относятся к номеру исходного файла и номеру страницы внутри него. PDF 1 — «Отсканированный документ.pdf»; PDF 2–4 — файлы с соответствующими цифрами в названии. Исходные листы, таблицы стоимости, расчёты и черновое распознавание сохранены в архиве сайта.')
p('Исторические цены и пищевая ценность приводятся как записи источника, без пересчёта и актуализации. Неясные и отсутствующие сведения не дополнены рецептами из интернета.')
doc.add_page_break();doc.add_heading('Содержание',1)
for cat in dict.fromkeys(r['category'] for r in R):
 head(cat)
 p(' · '.join(f"{r['number']:02}. {r['title']}" for r in R if r['category']==cat))
# Appendix list omitted here to keep the contents on one page.
for r in R:
 doc.add_page_break();p(f"{r['category']}  /  {r['number']:02}")
 doc.add_heading(r['title'],1)
 p('Источники: '+ '; '.join(f"PDF {s['pdf']}, с. {s['page']}" for s in r['sources']))
 if r['kind']=='fragment':p('Фрагмент рецепта с неполными количествами или технологией')
 if r['yieldText']:p('Выход: '+r['yieldText'])
 table([(i['name'],i['amount']) for i in r['ingredients']],['Ингредиент','Количество'])
 if r['steps']:
  head('Приготовление')
  for i,s in enumerate(r['steps'],1):p(f'{i}. {s}')
 if r['notes']:
  head('Варианты и записи на полях')
  for s in r['notes']:p(s)
 if r['issues']:
  head('Уточнения к источнику')
  for s in r['issues']:p(s)
doc.add_page_break();doc.add_heading('Кухонные заметки',1)
for n in N:
 note_start=len(doc.paragraphs)
 head(n['title']);p('Источники: '+ '; '.join(f"PDF {s['pdf']}, с. {s['page']}" for s in n['sources']))
 p(n['text'])
 if n['issue']:p('Уточнение: '+n['issue'])
 for paragraph in doc.paragraphs[note_start:-1]:paragraph.paragraph_format.keep_with_next=True
doc.add_page_break();doc.add_heading('Реестр исходных страниц',1)
registry=[]
for d,count in [(1,24),(2,24),(3,24),(4,22)]:
 if d>1:doc.add_page_break()
 head(f'PDF {d}')
 rows=[]
 for page in range(1,count+1):
  titles=[r['title'] for r in R if dict(pdf=d,page=page) in r['sources']]+[n['title'] for n in N if dict(pdf=d,page=page) in n['sources']]
  rows.append((page,'; '.join(titles)));registry.append(dict(pdf=d,page=page,entries=titles))
 table(rows,['Страница','Куда включена'])
doc.core_properties.title='Книга рецептов';doc.core_properties.subject='Оцифрованные рецепты и рукописные записи';doc.core_properties.author=''
doc.save(OUT/'recipe-book.docx')
md=['# Книга рецептов','', '73 рецептурные записи. 94 исходные страницы. Пары чисел в граммах — брутто/нетто. Неполные записи и разночтения отмечены явно.','']
for r in R:
 md += [f"## {r['number']:02}. {r['title']}",'',r['category'],'', 'Источники: '+ '; '.join(f"PDF {s['pdf']}, с. {s['page']}" for s in r['sources']),'']
 if r['kind']=='fragment':md+=['**Фрагмент рецепта**','']
 if r['yieldText']:md+=['Выход: '+r['yieldText'],'']
 md+=['| Ингредиент | Количество |','|---|---|']+[f"| {i['name']} | {i['amount']} |" for i in r['ingredients']]+['']
 md += [f'{i}. {s}' for i,s in enumerate(r['steps'],1)]+['']
 for name,key in [('Варианты и записи на полях','notes'),('Уточнения к источнику','issues')]:
  if r[key]:md+=['### '+name,'']+['- '+s for s in r[key]]+['']
md+=['## Кухонные заметки','']
for n in N:md += ['### '+n['title'],'',n['text'],'',n['issue'],'','Источники: '+ '; '.join(f"PDF {s['pdf']}, с. {s['page']}" for s in n['sources']),'']
md+=['## Реестр исходных страниц','']+[f"- PDF {p['pdf']}, с. {p['page']}: {'; '.join(p['entries'])}" for p in registry]
(OUT/'recipe-book.md').write_text('\n'.join(md))
(ROOT/'data/recipes.json').write_text(json.dumps(dict(recipes=R,notes=N,pages=registry),ensure_ascii=False,indent=2))
print('Saved DOCX, Markdown and JSON')
